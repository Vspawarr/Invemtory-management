import os
from pathlib import Path
from typing import List, Optional
import openpyxl
from openpyxl.workbook import Workbook
import logging

from .models import SAPRequest
from .utils import get_month_folder

logger = logging.getLogger(__name__)

class ExcelGenerator:
    """
    Handles the generation of Month-Level Excel files containing exactly one row
    per successfully processed request per date. Prevents duplicates.
    """

    HEADERS = [
        "KC",
        "EinkBeleg",
        "Materialnummer",
        "Materialkurztext",
        "Input Date",
        "Supplier",
        "Supplier Email",
        "Designer Note",
        "U",
        "Supplier Source"
    ]

    def __init__(self, base_output_dir: Path):
        self.base_output_dir = Path(base_output_dir)

    def _get_excel_path(self, date_str: str) -> Path:
        month_folder, date_folder_name = get_month_folder(date_str)
        month_path = self.base_output_dir / month_folder
        month_path.mkdir(parents=True, exist_ok=True)
        return month_path / f"{date_folder_name}.xlsx"

    def append_request(self, request: SAPRequest):
        """
        Appends a processed request to the correct Date Excel.
        Validates to ensure no duplicate rows (Idempotency).
        """
        if not request.assigned_kc:
            logger.error(f"Cannot append request {request.eink_beleg} without KC assignment.")
            return

        excel_path = self._get_excel_path(request.input_date)

        if excel_path.exists():
            wb = openpyxl.load_workbook(excel_path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.append(self.HEADERS)

        # Check for duplicates
        row_exists = False
        if ws.max_row > 1:
            for row in ws.iter_rows(min_row=2, values_only=True):
                # Row mapping: KC is 0, EinkBeleg is 1, Materialnummer is 2
                if row[1] == request.eink_beleg and row[2] == request.material_nummer:
                    row_exists = True
                    break

        if row_exists:
            logger.info(f"Request {request.eink_beleg}-{request.material_nummer} already in {excel_path.name}. Skipping append.")
            return

        # Append new row
        row_data = [
            request.assigned_kc,
            request.eink_beleg,
            request.material_nummer,
            request.material_kurztext or "",
            request.input_date,
            request.supplier or "",
            request.supplier_email or "",
            request.designer_note or "",
            request.u or "",
            request.supplier_source or ""
        ]
        ws.append(row_data)
        wb.save(excel_path)
        logger.info(f"Appended {request.assigned_kc} to {excel_path.name}")
